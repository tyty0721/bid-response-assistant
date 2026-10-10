"""Deterministic checks for agent-authored, evidence-linked procurement records."""
import argparse
import copy
import csv
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import unicodedata

STATUSES = {'supported', 'partial', 'missing_evidence', 'conflict', 'not_applicable', 'needs_confirmation'}
KINDS = {'qualification', 'substantive', 'scoring', 'contract', 'procedure'}
PROOFS = {'none', 'promise', 'statement', 'license', 'historical_contract', 'screenshot', 'live_demo',
          'textbook', 'copyright', 'platform_record'}


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save(path, value):
    # Exclusive create prevents a resumed run from overwriting human changes.
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)


def fingerprint(record):
    data = {k: v for k, v in record.items() if k not in {'approval', 'state'}}
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True,
                                    separators=(',', ':')).encode()).hexdigest()


def money(value):
    if value is None or isinstance(value, bool) or not isinstance(value, (str, int)):
        raise ValueError('Amounts must be integer yuan or decimal strings, never floats/null')
    try:
        amount = Decimal(value)
    except InvalidOperation as error:
        raise ValueError('Invalid decimal amount') from error
    if not amount.is_finite() or amount < 0 or amount >= Decimal('1000000000000'):
        raise ValueError('Amount must be finite, nonnegative and below one trillion yuan')
    if amount != amount.quantize(Decimal('.01')):
        raise ValueError('Amount must be nonnegative, finite and have at most 2 decimal places')
    return amount


def rmb_upper(value):
    amount = money(value)
    if amount >= Decimal('1000000000000'):
        raise ValueError('RMB uppercase conversion supports amounts below one trillion yuan')
    digits = '零壹贰叁肆伍陆柒捌玖'
    units = ['', '拾', '佰', '仟']
    whole, cents = divmod(int(amount * 100), 100)
    def group(number):
        result, zero = '', False
        for pos in range(3, -1, -1):
            digit = number // (10 ** pos) % 10
            if digit:
                result += ('零' if zero and result else '') + digits[digit] + units[pos]
                zero = False
            elif result:
                zero = True
        return result
    text, zero = '', False
    for idx, unit in ((2, '亿'), (1, '万'), (0, '')):
        number = whole // (10000 ** idx) % 10000
        if number:
            if text and (zero or number < 1000):
                text += '零'
            text += group(number) + unit
            zero = False
        elif text:
            zero = True
    text = (text or '零') + '元'
    jiao, fen = divmod(cents, 10)
    if not cents:
        return text + '整'
    if jiao:
        text += digits[jiao] + '角'
    if fen:
        text += ('零' if not jiao else '') + digits[fen] + '分'
    return text


def validate(record):
    if record.get('schema_version') != 1:
        raise ValueError('schema_version must be 1')
    project = record['project']
    for key in ('id', 'name', 'package', 'version', 'method', 'category', 'deadline', 'rules', 'sources'):
        if key not in project:
            raise ValueError(f'Missing project.{key}')
    for key in ('requirements', 'evidence', 'responses', 'reviewed_blocks'):
        if not isinstance(record[key], list):
            raise ValueError(f'{key} must be an array')
    for collection in ('requirements', 'evidence', 'responses'):
        ids = [r['id'] if collection != 'responses' else r['requirement_id'] for r in record[collection]]
        if len(set(ids)) != len(ids):
            raise ValueError(f'Duplicate IDs in {collection}')
    requirement_ids = {r['id'] for r in record['requirements']}
    evidence_ids = {e['id'] for e in record['evidence']}
    for requirement in record['requirements']:
        if requirement['kind'] not in KINDS or requirement['proof_type'] not in PROOFS:
            raise ValueError(f'Invalid requirement kind/proof: {requirement["id"]}')
        if not isinstance(requirement['mandatory'], bool) or not isinstance(requirement['response_required_now'], bool):
            raise ValueError('mandatory and response_required_now must be booleans')
        for key in ('text', 'phase', 'sources'):
            if not requirement.get(key):
                raise ValueError(f'Missing requirement {key}')
    for response in record['responses']:
        if response['requirement_id'] not in requirement_ids or response['status'] not in STATUSES:
            raise ValueError('Response refers to unknown requirement or invalid status')
        if any(e not in evidence_ids for e in response['evidence_ids']):
            raise ValueError('Response refers to unknown evidence')
        if response['status'] == 'not_applicable' and not response.get('applicability_reason'):
            raise ValueError('not_applicable requires applicability_reason')
    if project['deadline'] is not None:
        deadline = datetime.fromisoformat(project['deadline'])
        if deadline.tzinfo is None:
            raise ValueError('Deadline must include timezone')
    for key in ('budget', 'ceiling'):
        if project['rules'].get(key) is not None:
            money(project['rules'][key])
    for key in ('performance_deposit_rate', 'agency_fee_rate', 'low_price_ratio', 'price_deduction_rate'):
        value = project['rules'].get(key)
        if value is not None:
            rate = Decimal(str(value))
            if not rate.is_finite() or not 0 <= rate <= 1:
                raise ValueError(f'Invalid proportion {key}')
    days = project['rules'].get('min_validity_days')
    if days is not None and (type(days) is not int or days < 0):
        raise ValueError('min_validity_days must be a nonnegative integer')


def check(record, workspace, as_of):
    validate(record)
    workspace = Path(workspace).resolve()
    manifest, blocks = load(workspace / 'manifest.json'), load(workspace / 'source_blocks.json')
    sources = {s['id']: s for s in manifest['sources']}
    block_map = {b['id']: b for b in blocks}
    materials_hash = hashlib.sha256((workspace / 'manifest.json').read_bytes() +
                                    (workspace / 'source_blocks.json').read_bytes()).hexdigest()
    issues, rows, calculations = [], [], {}
    def issue(code, message, rid=None, severity='blocker'):
        issues.append(dict(code=code, severity=severity, requirement_id=rid, message=message))
    if manifest.get('source_blocks_sha256') != hashlib.sha256((workspace / 'source_blocks.json').read_bytes()).hexdigest():
        issue('CORPUS_CHANGED', 'Extracted text differs from the intake snapshot; ingest a new version')
    for source in sources.values():
        path = (workspace / source['path']).resolve()
        if workspace not in path.parents or not path.is_file() or path.is_symlink():
            issue('SOURCE_PATH', f'Invalid or missing local source {source["id"]}')
        elif hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
            issue('SOURCE_CHANGED', f'Hash differs for {source["id"]}')
        if source['parser'] == 'failed':
            issue('PARSER_FAILED', f'Source cannot be read: {source["id"]}')
    def citation(ref, rid=None):
        block = block_map.get(ref.get('block_id'))
        if not block or ref.get('source_id') != block['source_id'] or block['source_id'] not in sources:
            issue('CITATION_INVALID', 'Citation source or block does not exist', rid)
            return False
        quote = unicodedata.normalize('NFKC', ref.get('quote', '')).strip()
        if not quote or quote not in unicodedata.normalize('NFKC', block['text']):
            issue('CITATION_QUOTE', 'Quoted text does not occur in the cited block', rid)
            return False
        if block['quality'] != 'native':
            issue('PARSE_REVIEW', 'Cited block requires visual/OCR review', rid)
            return False
        return True
    project = record['project']
    for field in ('id', 'name', 'package', 'version', 'company'):
        if not project.get(field):
            issue('PROJECT_FIELD_MISSING', f'Project or supplier field is unknown: {field}')
    supported = (project['method'] == 'competitive_consultation' and project['category'] == 'services'
                 and project.get('jurisdiction') == 'CN-government-procurement')
    if not supported:
        issue('UNSUPPORTED_METHOD', 'Only limited source analysis is supported for this procurement method')
    if project['deadline'] is None:
        issue('DEADLINE_MISSING', 'Supply the official notice and deadline; do not infer a date from the version ID')
    elif datetime.fromisoformat(project['deadline']).date() < as_of:
        issue('DEADLINE_PASSED', 'Deadline is before the review date; use historical analysis only')
    if not record['requirements']:
        issue('REQUIREMENTS_EMPTY', 'No requirements were extracted; a review cannot be complete')
    for ref in project['sources']:
        citation(ref)
    if not project['sources']:
        issue('PROJECT_SOURCE', 'Project facts need procurement source citations')
    procurement_sources = set(record.get('procurement_source_ids', []))
    expected_blocks = {b['id'] for b in blocks if b['source_id'] in procurement_sources}
    reviewed = set(record['reviewed_blocks'])
    if not procurement_sources or not procurement_sources.issubset(sources):
        issue('PROCUREMENT_SCOPE', 'Explicitly identify procurement sources')
    for block_id in sorted(expected_blocks - reviewed):
        issue('COVERAGE_UNREVIEWED', f'No recorded semantic review for {block_id}')
    if reviewed - block_map.keys():
        issue('REVIEW_BLOCK_INVALID', 'Review ledger contains nonexistent blocks')
    responses = {r['requirement_id']: r for r in record['responses']}
    evidence = {e['id']: e for e in record['evidence']}
    for requirement in record['requirements']:
        rid = requirement['id']
        if requirement.get('package') not in (None, project['package']):
            issue('PACKAGE_MISMATCH', 'Requirement belongs to a different procurement package', rid)
        for ref in requirement['sources']:
            citation(ref, rid)
        response = responses.get(rid)
        severity = 'blocker' if requirement['mandatory'] else 'warning'
        status = response['status'] if response else 'missing_evidence'
        text = response.get('text') if response else None
        if not requirement['response_required_now'] and not response:
            status = 'reference_only'
            text = '记录项目规则；当前无需为本条单独提交响应或证明。'
        if requirement['response_required_now'] and (not response or status not in {'supported', 'not_applicable'}):
            issue('RESPONSE_GAP', 'Missing, partial, conflicting or unconfirmed response', rid, severity)
        if response and status == 'supported':
            if not text or not response['evidence_ids']:
                issue('UNSUPPORTED_CLAIM', 'supported requires response text and appropriate evidence', rid, severity)
                status = 'missing_evidence'
            for eid in response['evidence_ids']:
                e = evidence[eid]
                good = True
                if e.get('type') != requirement['proof_type']:
                    issue('EVIDENCE_TYPE', f'{eid} is not {requirement["proof_type"]}', rid, severity)
                    good = False
                if e.get('entity') != project.get('company') or not project.get('company'):
                    issue('ENTITY_MISMATCH', f'{eid} does not match the supplier entity', rid, severity)
                    good = False
                if not e.get('confirmed_by') or not e.get('confirmed_at'):
                    issue('EVIDENCE_UNCONFIRMED', f'{eid} needs accountable human confirmation', rid, severity)
                    good = False
                if e.get('valid_until') and date.fromisoformat(e['valid_until']) < as_of:
                    issue('EVIDENCE_EXPIRED', f'{eid} expired before the review date', rid, severity)
                    good = False
                if requirement.get('product_version') and e.get('product_version') != requirement['product_version']:
                    issue('PRODUCT_MISMATCH', f'{eid} is for a different product version', rid, severity)
                    good = False
                if e.get('type') not in {'promise', 'none'} and not e.get('sources'):
                    issue('EVIDENCE_SOURCE', f'{eid} needs an original proof source', rid, severity)
                    good = False
                for ref in e.get('sources', []):
                    good = citation(ref, rid) and good
                    if ref.get('source_id') in procurement_sources:
                        issue('SELF_PROOF', 'Procurement requirements cannot prove supplier capability', rid, severity)
                        good = False
                if e.get('type') == 'promise' and not e.get('statement'):
                    issue('PROMISE_EMPTY', 'A confirmed promise must record the actual commitment', rid, severity)
                    good = False
                if not good:
                    status = 'needs_confirmation'
        if response and status == 'not_applicable' and not response.get('confirmed_by'):
            issue('APPLICABILITY_UNCONFIRMED', 'Conditional exemption needs a human-confirmed reason', rid, severity)
            status = 'needs_confirmation'
        for forbidden in record.get('review_terms', {}).get('wrong_project_names', []):
            if text and forbidden in text:
                issue('WRONG_PROJECT', f'Response contains another project/entity: {forbidden}', rid, severity)
        rows.append({'id': rid, 'kind': requirement['kind'], 'text': requirement['text'],
                     'status': status, 'response': text, 'evidence_ids': response['evidence_ids'] if response else [],
                     'sources': requirement['sources']})
    rules, quote = project['rules'], record.get('quote')
    if not quote or quote.get('total') is None:
        issue('QUOTE_MISSING', 'Supplier price is unknown; the procurement ceiling is not a quote')
    else:
        try:
            total = money(quote['total'])
            if quote.get('currency') != 'CNY' or quote.get('tax_inclusive') is not True:
                issue('QUOTE_BASIS', 'Currency or tax basis is unknown or differs; confirm the documented project basis')
            if quote.get('project_version') != project['version']:
                issue('QUOTE_STALE', 'Quote uses another procurement version')
            if quote.get('round') == 'final' and not supported:
                issue('FINAL_QUOTE_UNSUPPORTED', 'Do not apply consultation final quotation to this method')
            if quote.get('uppercase') != rmb_upper(total.to_eng_string()):
                issue('QUOTE_UPPERCASE', 'Uppercase and numeric quote differ; do not silently correct an approved quote')
            if not quote.get('approved_by') or not quote.get('approved_at'):
                issue('QUOTE_UNAPPROVED', 'Pricing decision needs company approval')
            ceiling = money(rules['ceiling']) if rules.get('ceiling') is not None else None
            if ceiling is None:
                issue('CEILING_MISSING', 'An unknown ceiling prevents a completed quote check')
            elif total > ceiling:
                issue('QUOTE_CEILING', 'Quote exceeds the package ceiling')
            if quote.get('items'):
                summed = sum((money(i['unit_price']) * Decimal(str(i['quantity'])) for i in quote['items']), Decimal(0))
                if not summed.is_finite() or any(Decimal(str(i['quantity'])) <= 0 for i in quote['items']):
                    raise ValueError('Quantities must be finite and positive')
                if summed.quantize(Decimal('.01'), rounding=ROUND_HALF_UP) != total:
                    issue('QUOTE_SUM', 'Line-item total differs from the declared quote')
            if rules.get('low_price_ratio') is not None and ceiling is not None:
                boundary = ceiling * Decimal(str(rules['low_price_ratio']))
                if total < boundary:
                    issue('LOW_PRICE_REVIEW', 'Below the documented threshold; prepare cost explanation, not automatic invalidation', severity='warning')
            for rule, output in (('performance_deposit_rate', 'performance_deposit'), ('agency_fee_rate', 'agency_fee')):
                if rules.get(rule) is not None:
                    calculations[output] = str((total * Decimal(str(rules[rule]))).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))
            calculations['response_deposit'] = rules.get('response_deposit')
            calculations['actual_price_score'] = None
            if quote.get('price_deduction_rate', '0') != str(rules.get('price_deduction_rate', '0')):
                issue('PRICE_DEDUCTION', 'Quoted price deduction is not supported by this project')
            if rules.get('min_validity_days') is None:
                issue('VALIDITY_RULE_MISSING', 'Documented validity requirement is unknown; do not infer a minimum')
            elif type(quote.get('validity_days')) is not int or quote['validity_days'] < rules['min_validity_days']:
                issue('VALIDITY_SHORT', 'Response validity is missing or shorter than the documented minimum')
        except (ValueError, InvalidOperation, KeyError, TypeError) as error:
            issue('QUOTE_INVALID', str(error))
    for conflict in record.get('conflicts', []):
        for ref in conflict.get('sources', []):
            citation(ref)
        if not conflict.get('resolution'):
            issue('DOCUMENT_CONFLICT', conflict['description'])
    scoring = record.get('scoring', {})
    counts = Counter(r.get('marker', 'unmarked') for r in record['requirements']
                     if r.get('scoring_group') == 'technical')
    for marker, expected in scoring.get('technical_expected_counts', {}).items():
        if counts[marker] != expected:
            issue('SCORING_COUNT_MISMATCH', f'{marker}: extracted {counts[marker]}, scoring table says {expected}')
    calculations['technical_marker_counts'] = dict(counts)
    approval = record.get('approval')
    if approval and (approval.get('input_hash') != fingerprint(record) or approval.get('materials_hash') != materials_hash):
        issue('APPROVAL_STALE', 'Inputs changed after approval; obtain a fresh approval')
    state = 'needs_information' if any(i['severity'] == 'blocker' for i in issues) else 'reviewed'
    if state == 'reviewed' and approval and approval.get('reviewer') and approval.get('manual_checklist'):
        state = 'approved_for_manual_submission'
    return {'schema_version': 1, 'project_id': project['id'], 'input_hash': fingerprint(record),
            'materials_hash': materials_hash,
            'as_of': as_of.isoformat(), 'state': state, 'issues': issues, 'matrix': rows,
            'calculations': calculations, 'semantic_review': 'agent_and_human_required',
            'submission': 'not_performed'}


def differences(old, new):
    old_rows = {r['id']: r for r in old['requirements']}
    new_rows = {r['id']: r for r in new['requirements']}
    changed = sorted(k for k in old_rows.keys() | new_rows.keys() if old_rows.get(k) != new_rows.get(k))
    project_changed = old['project'] != new['project']
    other_changed = any(old.get(k) != new.get(k) for k in old.keys() | new.keys()
                        if k not in {'requirements', 'project', 'approval', 'state'})
    revised = copy.deepcopy(new)
    if changed or project_changed or other_changed:
        revised['approval'] = None
        revised['state'] = 'needs_information'
    return {'changed_requirements': changed, 'project_changed': project_changed,
            'other_inputs_changed': other_changed,
            'affected_sections': sorted({r.get('section', r['id']) for r in old['requirements'] + new['requirements']
                                         if r['id'] in changed or project_changed or other_changed}),
            'quote_review_required': bool(project_changed or other_changed or changed),
            'approval_invalidated': bool(changed or project_changed or other_changed), 'updated_record': revised}


def export(record, checks, output):
    if checks['input_hash'] != fingerprint(record):
        raise ValueError('Checks belong to a different input version')
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.3)
    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[style_name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    document.styles['Normal'].font.size = Pt(10.5)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.add_paragraph('供应商响应文件审核稿', 'Title')
    document.add_paragraph(record['project']['name'])
    document.add_paragraph(f'项目 {record["project"]["id"]}  采购包 {record["project"]["package"]}  版本 {record["project"]["version"]}')
    document.add_paragraph(f'状态 {checks["state"]}。这是内部审核草稿。待补内容必须经企业确认并按正式平台模板完成签章和递交。')
    document.add_heading('项目与报价', level=1)
    supplier = record['project'].get('company') or '【待补企业主体】'
    total = (record.get('quote') or {}).get('total')
    document.add_paragraph(f'供应商 {supplier}。响应截止时间 {record["project"]["deadline"] or "【待官方公告确认】"}。')
    document.add_paragraph(f'报价 {total if total is not None else "【待企业批准报价】"}。预算及限价不能代替企业报价。')
    document.add_heading('逐项响应', level=1)
    for row in checks['matrix']:
        document.add_heading(row['id'], level=2)
        document.add_paragraph(row['text'])
        document.add_paragraph(row['response'] or '【待补证据或企业确认后起草响应】')
        citations = '；'.join(f'{s["source_id"]} {s["block_id"]}' for s in row['sources'])
        document.add_paragraph(f'状态 {row["status"]}。依据 {citations}。证据 {", ".join(row["evidence_ids"]) or "尚未提供"}。')
    document.add_heading('递交前人工操作', level=1)
    for line in ('核对正式公告和全部更正，确认获取凭证及截止时间。',
                 '复核主体、声明、报价、附件、履约承诺与真实系统演示。',
                 '按平台指定格式制作并由授权人员签章、加密、上传和核验回执。',
                 '本技能没有执行报名、签章、递交、解密或联系采购代理机构。'):
        document.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.add_run('内部审核稿  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.save(output / 'response-draft.docx')
    with (output / 'response-matrix.csv').open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['要求ID', '类型', '原文', '状态', '响应', '证据', '出处'])
        for row in checks['matrix']:
            values = [row['id'], row['kind'], row['text'], row['status'], row['response'] or '待补',
                      ','.join(row['evidence_ids']), ';'.join(s['block_id'] for s in row['sources'])]
            writer.writerow(["'" + v if v.lstrip().startswith(('=', '+', '-', '@')) else v for v in values])
    lines = ['# 响应审查报告', '', f'项目：{checks["project_id"]}；状态：{checks["state"]}', '',
             '程序检查不能证明证据真实或保证合规；需复核全文语义和正式公告。', '',
             '| 等级 | 问题 | 要求 | 处理 |', '|---|---|---|---|']
    lines += [f'| {i["severity"]} | {i["code"]} | {i["requirement_id"] or "项目"} | {i["message"].replace(chr(124), chr(47)).replace(chr(10), " ")} |' for i in checks['issues']]
    (output / 'review_report.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    save(output / 'checks.json', checks)
    save(output / 'evidence-index.json', record['evidence'])
    save(output / 'handoff_manifest.json', {'project_id': checks['project_id'], 'state': checks['state'],
         'input_hash': checks['input_hash'], 'materials_hash': checks['materials_hash'], 'submission': 'not_performed',
         'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in output.iterdir() if p.is_file()}})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    for name in ('validate', 'check', 'export', 'approve'):
        p = commands.add_parser(name)
        p.add_argument('record')
        if name != 'validate':
            p.add_argument('--workspace', required=True)
            p.add_argument('--as-of', required=True, help='ISO review date; no implicit machine clock')
            p.add_argument('--output', required=True)
        if name == 'approve':
            p.add_argument('--reviewer', required=True)
            p.add_argument('--manual-checklist', required=True, help='Actual human review record, not generated approval')
    p = commands.add_parser('diff')
    p.add_argument('old')
    p.add_argument('new')
    p.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        if args.command == 'diff':
            save(args.output, differences(load(args.old), load(args.new)))
            return
        record = load(args.record)
        validate(record)
        if args.command == 'validate':
            print('Record structure valid; facts and semantics are not proven')
            return
        checks = check(record, args.workspace, date.fromisoformat(args.as_of))
        if args.command == 'export':
            export(record, checks, args.output)
        elif args.command == 'approve':
            if checks['state'] == 'needs_information' or checks['issues']:
                raise ValueError('Resolve blockers and warnings before recording final manual approval')
            record['approval'] = {'reviewer': args.reviewer, 'manual_checklist': args.manual_checklist,
                                  'at': datetime.now(timezone.utc).isoformat(), 'input_hash': fingerprint(record),
                                  'materials_hash': checks['materials_hash']}
            record['state'] = 'approved_for_manual_submission'
            save(args.output, record)
        else:
            save(args.output, checks)
        print(json.dumps({'state': checks['state'], 'issues': len(checks['issues']), 'submission': 'not_performed'}))
    except (ValueError, KeyError, TypeError, OSError, InvalidOperation) as error:
        parser.exit(2, f'{error}\n')


if __name__ == '__main__':
    main()
